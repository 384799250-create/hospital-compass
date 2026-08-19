"""Import the user-provided medical-insurance workbook without personal contacts."""

from datetime import UTC, datetime
import json
from pathlib import Path
import sys
import time

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from app.excel_importer import record_from_row
from app.hospital_store import upsert_hospitals_bulk

DATA_DIR = Path(r'F:\hospital-compass-data')
SOURCE_FILE = next(Path(r'D:\BaiduNetdiskDownload').rglob('【EXCEL数据】全国医院数据(医疗保险)/*.xlsx'))
SOURCE_URL = 'file:///' + SOURCE_FILE.as_posix()


def main() -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    progress_path = DATA_DIR / 'medical-insurance-import-progress.jsonl'
    completed = _completed(progress_path)
    workbook = load_workbook(SOURCE_FILE, read_only=True, data_only=True)
    sheet = workbook.active
    imported = 0
    batch: list[tuple[dict[str, object], list[dict[str, str]]]] = []
    for index, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
        if index in completed:
            continue
        record = record_from_row(row)
        try:
            if record['name']:
                batch.append((record, [{
                    'source_url': SOURCE_URL,
                    'source_type': '用户提供医保医院数据',
                    'title': SOURCE_FILE.name,
                }]))
                if len(batch) >= 500:
                    imported += upsert_hospitals_bulk(batch)
                    batch.clear()
            _log(progress_path, {'row': index, 'status': 'completed'})
        except Exception as error:
            _log(progress_path, {'row': index, 'status': 'failed', 'error': type(error).__name__})
        if index % 500 == 0:
            time.sleep(0.1)
    if batch:
        imported += upsert_hospitals_bulk(batch)
    _log(progress_path, {'status': 'run_completed', 'imported': imported, 'rows': sheet.max_row - 1})


def _completed(path: Path) -> set[int]:
    if not path.exists():
        return set()
    rows = set()
    for line in path.read_text(encoding='utf-8').splitlines():
        try:
            record = json.loads(line)
            if record.get('status') == 'completed' and 'row' in record:
                rows.add(int(record['row']))
        except (ValueError, KeyError, json.JSONDecodeError):
            continue
    return rows


def _log(path: Path, record: dict[str, object]) -> None:
    with path.open('a', encoding='utf-8') as file:
        file.write(json.dumps({'at': datetime.now(UTC).isoformat(), **record}, ensure_ascii=False) + '\n')


if __name__ == '__main__':
    main()
